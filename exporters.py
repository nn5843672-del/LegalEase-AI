from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def export_file(file_type: str, title: str, content: str) -> tuple[bytes, str, str]:
    safe_name = "legalease-draft"
    if file_type == "txt":
        return content.encode("utf-8"), "text/plain; charset=utf-8", f"{safe_name}.txt"

    if file_type == "docx":
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        normal = document.styles["Normal"]
        normal.font.name = "Aptos"
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGBColor(35, 49, 43)
        document.add_paragraph("LEGALEASE  /  DOCUMENT STUDIO", style="Subtitle")
        heading = document.add_heading(title or "Legal document", 0)
        heading.style.font.color.rgb = RGBColor(18, 91, 69)
        for line in content.splitlines():
            line = line.strip()
            if not line:
                document.add_paragraph()
            elif line.isupper() and len(line) < 90:
                document.add_heading(line.title(), level=2)
            else:
                document.add_paragraph(line)
        footer = section.footer.paragraphs[0]
        footer.alignment = 2
        footer.add_run("LegalEase • Draft for review").font.color.rgb = RGBColor(92, 111, 102)
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", f"{safe_name}.docx"

    if file_type == "pdf":
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=22*mm, leftMargin=22*mm,
                                topMargin=25*mm, bottomMargin=22*mm, title=title)
        base = getSampleStyleSheet()
        heading_style = ParagraphStyle("DocHeading", parent=base["Heading2"], textColor=colors.HexColor("#155b45"),
                                       fontName="Helvetica-Bold", fontSize=10.5, leading=15, spaceBefore=9, spaceAfter=4)
        body = ParagraphStyle("DocBody", parent=base["BodyText"], fontName="Helvetica", fontSize=9.5,
                              leading=14, textColor=colors.HexColor("#26332e"), alignment=TA_LEFT)
        title_style = ParagraphStyle("DocTitle", parent=base["Title"], fontName="Helvetica-Bold", fontSize=20,
                                     leading=24, textColor=colors.HexColor("#155b45"), alignment=TA_CENTER, spaceAfter=14)
        story = [Paragraph("LEGALEASE  /  DOCUMENT STUDIO", ParagraphStyle("Brand", parent=body, alignment=TA_CENTER,
                                                                            textColor=colors.HexColor("#74847c"))),
                 Spacer(1, 8), Paragraph(escape(title or "Legal document"), title_style)]
        for line in content.splitlines():
            line = line.strip()
            if not line:
                story.append(Spacer(1, 5))
            elif line.isupper() and len(line) < 90:
                story.append(Paragraph(escape(line.title()), heading_style))
            else:
                story.append(Paragraph(escape(line).replace("\t", "&nbsp;&nbsp;"), body))
        def page_footer(canvas, _doc):
            canvas.saveState()
            canvas.setStrokeColor(colors.HexColor("#e5ebe7"))
            canvas.line(22*mm, 15*mm, A4[0]-22*mm, 15*mm)
            canvas.setFont("Helvetica", 8)
            canvas.setFillColor(colors.HexColor("#74847c"))
            canvas.drawString(22*mm, 10*mm, "LegalEase  •  Draft for review")
            canvas.drawRightString(A4[0]-22*mm, 10*mm, f"Page {_doc.page}")
            canvas.restoreState()
        doc.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
        return buffer.getvalue(), "application/pdf", f"{safe_name}.pdf"
    raise ValueError("Unsupported export format")
